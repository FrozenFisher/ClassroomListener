import sys
import pandas as pd
from datetime import datetime
from PyQt6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
                            QTableWidget, QTableWidgetItem, QLabel, QFileDialog,
                            QPushButton, QMessageBox, QAbstractScrollArea, QLayout,
                            QSizePolicy, QHeaderView)
from PyQt6.QtCore import Qt, QTimer
from PyQt6.QtGui import QColor
import os
from openpyxl import load_workbook
from PyQt6.QtWidgets import QHBoxLayout, QSpinBox
import sys as _sys

class ScheduleWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle('课程表显示')
        self.setGeometry(100, 100, 800, 600)
        
        # 创建主窗口部件
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        layout = QVBoxLayout(central_widget)
        

        # 顶部控制区
        top_row = QHBoxLayout()
        self.file_button = QPushButton('选择课程表文件')
        self.file_button.clicked.connect(self.select_file)
        top_row.addWidget(self.file_button)

        # 字号调节
        self.font_label = QLabel('字号:')
        self.font_spin = QSpinBox()
        self.font_spin.setRange(8, 48)
        self.font_spin.setValue(14)
        self.font_spin.valueChanged.connect(self.apply_font_size)
        top_row.addWidget(self.font_label)
        top_row.addWidget(self.font_spin)

        # 保存设置
        self.save_button = QPushButton('保存设置')
        self.save_button.clicked.connect(self.save_settings_to_excel)
        top_row.addWidget(self.save_button)
        top_row.addStretch(1)
        layout.addLayout(top_row)
        
        # 创建日期和时间标签
        self.date_label = QLabel()
        self.time_label = QLabel()
        layout.addWidget(self.date_label)
        layout.addWidget(self.time_label)
        
        # 保存当前文件路径
        self.current_file = None
        
        # 创建课程表格
        self.table = QTableWidget()
        self.table.setColumnCount(2)
        self.table.setHorizontalHeaderLabels(['时间', '课程'])
        layout.addWidget(self.table)
        
        # 设置定时器更新时间
        self.timer = QTimer()
        self.timer.timeout.connect(self.update_time)
        self.timer.start(1000)  # 每秒更新一次
        
        # 初始化数据
        self.schedule_data = None
        self.duty_data = None  # 值日生列表 [(label, name)]
        self.load_schedule()
        # 读取设置并应用
        self.load_settings_from_excel()
        self.update_time()

        # 允许窗口与表格可伸缩
        self.table.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)
        
    def select_file(self):
        start_dir = self.get_app_dir()
        file_name, _ = QFileDialog.getOpenFileName(
            self,
            "选择课程表文件",
            start_dir,
            "Excel Files (*.xlsx *.xls);;All Files (*)"
        )
        if file_name:
            self.current_file = self.resolve_path(file_name)
            self.load_schedule()
            
    def load_schedule(self):
        try:
            # 如果没有选择文件，尝试从默认位置读取
            if not self.current_file:
                default_file = os.path.join(self.get_app_dir(), 'schedule.xlsx')
                if os.path.exists(default_file):
                    self.current_file = default_file
                else:
                    QMessageBox.warning(self, "提示", "未找到默认课程表文件，请选择文件")
                    return
            
            # 读取Excel文件
            df = pd.read_excel(self.current_file)
            weekday = datetime.now().weekday()  # 0-6，0是周一
            
            # 获取对应星期的课程（第1列是周一，第3列是周二，以此类推）
            time_col = weekday * 2
            course_col = time_col + 1
            
            # 提取当天的时间和课程：遇到 END_C 停止课程读取；之后到 END 为值日生
            raw_times = df.iloc[1:, time_col].tolist()
            raw_courses = df.iloc[1:, course_col].tolist()
            class_rows = []
            duty_rows = []
            mode = 'class'
            for t, c in zip(raw_times, raw_courses):
                t_str = '' if pd.isna(t) else str(t).strip()
                c_str = '' if pd.isna(c) else str(c).strip()
                token = t_str or c_str
                if token == 'END_C':
                    mode = 'duty'
                    continue
                if token == 'END':
                    break
                if mode == 'class':
                    class_rows.append((t_str, c_str))
                else:
                    duty_rows.append((t_str, c_str))

            times = [t for t, _ in class_rows]
            courses = [c for _, c in class_rows]
            
            # 更新表格
            self.table.setRowCount(len(times) + len(duty_rows))
            # 填充课程行
            row_index = 0
            for time, course in zip(times, courses):
                self.table.setItem(row_index, 0, QTableWidgetItem(time))
                self.table.setItem(row_index, 1, QTableWidgetItem(course))
                row_index += 1
            # 填充值日生行（追加）
            for label, name in duty_rows:
                self.table.setItem(row_index, 0, QTableWidgetItem(label))
                self.table.setItem(row_index, 1, QTableWidgetItem(name))
                row_index += 1
            
            # 调整表格尺寸以适配内容（初始）
            self.adjust_table_to_contents()
            # 如有保存的列宽，覆盖自适应结果
            self.apply_saved_column_widths()
            
            self.schedule_data = list(zip(times, courses))
            self.duty_data = duty_rows
        except Exception as e:
            print(f"加载课程表出错: {e}")
    
    def update_time(self):
        current_time = datetime.now()
        weekdays = ['一', '二', '三', '四', '五', '六', '日']
        
        # 更新日期和时间标签
        self.date_label.setText(f'星期{weekdays[current_time.weekday()]}')
        self.time_label.setText(current_time.strftime('%H:%M:%S'))
        
        if self.schedule_data:
            self.highlight_next_class(current_time)
            self.highlight_duty_period(current_time)
    
    def highlight_next_class(self, current_time):
        # 优先查找当前正在进行的课程；若无，则回退到下一节未开始的课程
        now_time = current_time.time()
        current_class_index = -1
        next_class_index = -1

        def parse_hhmm_to_time(hhmm: str):
            """将形如 H:MM 或 HH:MM 的时间字符串解析为 time 对象。"""
            hhmm = hhmm.strip()
            try:
                # 兼容 '8:00' / '08:00'
                parts = hhmm.split(':')
                if len(parts) != 2:
                    return None
                h = parts[0].zfill(2)
                m = parts[1].zfill(2)
                return datetime.strptime(f"{h}:{m}", '%H:%M').time()
            except Exception:
                return None

        for i, (time_str, _) in enumerate(self.schedule_data):
            if pd.isna(time_str):
                continue
            try:
                # 预期格式: "HH:MM-HH:MM" 或 "H:MM-H:MM"
                start_end = str(time_str).split('-')
                if len(start_end) != 2:
                    continue
                start_t = parse_hhmm_to_time(start_end[0])
                end_t = parse_hhmm_to_time(start_end[1])
                if not start_t or not end_t:
                    continue

                # 当前课：start <= now < end
                if start_t <= now_time < end_t:
                    current_class_index = i
                    break

                # 记录下一节（第一条 start > now 的课）
                if next_class_index == -1 and start_t > now_time:
                    next_class_index = i
            except Exception:
                continue
        
        # 重置所有行的背景色
        for i in range(self.table.rowCount()):
            for j in range(2):
                item = self.table.item(i, j)
                if item:
                    item.setBackground(QColor(255, 255, 255))
        
        # 优先高亮当前课；若没有当前课则高亮下一节
        target_index = current_class_index if current_class_index != -1 else next_class_index
        if target_index != -1:
            for j in range(2):
                item = self.table.item(target_index, j)
                if item:
                    item.setBackground(QColor(255, 255, 0))

    def highlight_duty_period(self, current_time):
        if not self.duty_data:
            return
        hour = current_time.hour
        # 上午 5:00-12:00 高亮以“上午”开头的值日生行；下午 12:00-22:00 高亮以“下午”开头
        highlight_prefix = None
        if 5 <= hour < 12:
            highlight_prefix = '上午'
        elif 12 <= hour < 22:
            highlight_prefix = '下午'
        if not highlight_prefix:
            return

        # 课程行数量用于偏移
        class_count = len(self.schedule_data) if self.schedule_data else 0
        for idx, (label, _) in enumerate(self.duty_data):
            if isinstance(label, str) and label.startswith(highlight_prefix):
                table_row = class_count + idx
                # 同时高亮时间列与姓名列
                for j in range(2):
                    item = self.table.item(table_row, j)
                    if item:
                        item.setBackground(QColor(255, 255, 0))

    def adjust_table_to_contents(self):
        # 根据内容自适应列、行尺寸，但不固定大小，允许拖动扩展
        self.table.resizeColumnsToContents()
        self.table.resizeRowsToContents()
        self.table.setSizeAdjustPolicy(QAbstractScrollArea.SizeAdjustPolicy.AdjustToContents)
        self.table.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.table.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.table.updateGeometry()

    # ------------------ 设置读写 ------------------
    def load_settings_from_excel(self):
        if not self.current_file:
            return
        try:
            if not self.current_file.lower().endswith('.xlsx'):
                return
            wb = load_workbook(self.current_file)
            ws = wb.active
            settings = {}
            # 从 A30 开始读取，最多读取到 A400 以防过大
            for row in range(30, 401):
                key = ws[f'A{row}'].value
                val = ws[f'B{row}'].value
                if key is None and val is None:
                    continue
                if isinstance(key, str):
                    settings[key.strip()] = val

            # 保存到实例，供后续使用
            self._loaded_settings = settings

            # 应用字号
            if 'font_size' in settings and settings['font_size']:
                try:
                    size = int(settings['font_size'])
                    if 8 <= size <= 48:
                        self.font_spin.setValue(size)
                        self.apply_font_size()
                except Exception:
                    pass

            # 应用窗口大小
            if 'window_width' in settings and 'window_height' in settings:
                try:
                    w = int(settings['window_width'])
                    h = int(settings['window_height'])
                    if w > 200 and h > 200:
                        self.resize(w, h)
                except Exception:
                    pass

            # 如有 excel_path，尝试切换并重新加载（支持相对路径：相对应用目录）
            if 'excel_path' in settings and isinstance(settings['excel_path'], str):
                p = self.resolve_path(settings['excel_path'])
                if os.path.isfile(p) and p.lower().endswith(('.xlsx', '.xls')):
                    if self.current_file != p:
                        self.current_file = p
                        self.load_schedule()
        except Exception as e:
            print(f"读取设置失败: {e}")

    def save_settings_to_excel(self, show_message=True):
        if not self.current_file:
            if show_message:
                QMessageBox.information(self, '提示', '请先选择课程表文件再保存设置')
            return
        if not self.current_file.lower().endswith('.xlsx'):
            if show_message:
                QMessageBox.warning(self, '提示', '仅支持将设置写入 .xlsx 文件')
            return
        try:
            # 保存到与 current_file 指定的路径；若相对路径则已被 resolve_path 处理
            wb = load_workbook(self.current_file)
            ws = wb.active

            # 将键写入 A 列，值写入 B 列，从 A30 起
            kv = {
                'font_size': self.font_spin.value(),
                'window_width': self.width(),
                'window_height': self.height(),
                'excel_path': self.current_file or '',
                'col_width_0': self.table.columnWidth(0) if self.table.columnCount() > 0 else 0,
                'col_width_1': self.table.columnWidth(1) if self.table.columnCount() > 1 else 0,
            }

            # 先构建现有键索引
            key_to_row = {}
            first_empty = None
            for row in range(30, 401):
                key_cell = ws[f'A{row}']
                val_cell = ws[f'B{row}']
                if key_cell.value is None and val_cell.value is None and first_empty is None:
                    first_empty = row
                if isinstance(key_cell.value, str):
                    key_to_row[key_cell.value.strip()] = row

            # 写回/追加
            for key, val in kv.items():
                if key in key_to_row:
                    r = key_to_row[key]
                else:
                    r = first_empty if first_empty is not None else 30
                    # 若占用则向下寻找空行
                    while ws[f'A{r}'].value is not None or ws[f'B{r}'].value is not None:
                        r += 1
                ws[f'A{r}'] = key
                ws[f'B{r}'] = val

            wb.save(self.current_file)
            if show_message:
                QMessageBox.information(self, '提示', '设置已保存到 Excel')
        except Exception as e:
            QMessageBox.critical(self, '错误', f'保存设置失败: {e}')

    def apply_font_size(self):
        size = self.font_spin.value()
        font = self.table.font()
        font.setPointSize(size)
        self.table.setFont(font)
        # 同步标签字号
        f2 = self.date_label.font()
        f2.setPointSize(size)
        self.date_label.setFont(f2)
        f3 = self.time_label.font()
        f3.setPointSize(size)
        self.time_label.setFont(f3)
        # 调整内容行高列宽
        self.adjust_table_to_contents()
        # 字号变化后再次应用可能保存的列宽
        self.apply_saved_column_widths()

    def closeEvent(self, event):
        # 关闭时自动尝试保存设置（若是 .xlsx）
        try:
            if self.current_file and self.current_file.lower().endswith('.xlsx'):
                self.save_settings_to_excel(show_message=False)
        finally:
            super().closeEvent(event)

    def apply_saved_column_widths(self):
        settings = getattr(self, '_loaded_settings', {}) or {}
        try:
            w0 = int(settings.get('col_width_0')) if settings.get('col_width_0') is not None else None
            w1 = int(settings.get('col_width_1')) if settings.get('col_width_1') is not None else None
            if w0 and self.table.columnCount() > 0:
                self.table.setColumnWidth(0, max(10, w0))
            if w1 and self.table.columnCount() > 1:
                self.table.setColumnWidth(1, max(10, w1))
        except Exception:
            pass

    # ------------------ 路径工具 ------------------
    def get_app_dir(self):
        try:
            if getattr(_sys, 'frozen', False) and hasattr(_sys, '_MEIPASS'):
                return os.path.dirname(_sys.executable)
            # 脚本运行
            return os.path.dirname(os.path.abspath(__file__))
        except Exception:
            return os.getcwd()

    def resolve_path(self, path_str):
        try:
            if not path_str:
                return path_str
            # 已是绝对路径
            if os.path.isabs(path_str):
                return os.path.normpath(path_str)
            # 相对路径：基于应用目录
            return os.path.normpath(os.path.join(self.get_app_dir(), path_str))
        except Exception:
            return path_str

def main():
    app = QApplication(sys.argv)
    window = ScheduleWindow()
    window.show()
    sys.exit(app.exec())

if __name__ == '__main__':
    main()

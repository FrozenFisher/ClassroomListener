import sys
import pandas as pd
from datetime import datetime, date, time
from PyQt6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
                            QTableWidget, QTableWidgetItem, QLabel, QFileDialog,
                            QPushButton, QMessageBox, QHBoxLayout, QCheckBox, QTimeEdit,
                            QAbstractScrollArea, QLayout, QSizePolicy, QHeaderView, QSpinBox)
from PyQt6.QtCore import Qt, QTimer
from PyQt6.QtGui import QColor
import os
from openpyxl import load_workbook


class ScheduleWindowTest(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle('课程表显示（测试：可手动时间）')
        self.setGeometry(120, 120, 820, 640)

        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        layout = QVBoxLayout(central_widget)

        # 顶部控制区
        top_row = QHBoxLayout()
        self.file_button = QPushButton('选择课程表文件')
        self.file_button.clicked.connect(self.select_file)
        top_row.addWidget(self.file_button)

        # 字号调节
        self.font_label = QLabel('字号:')
        self.font_spin = QSpinBox()
        self.font_spin.setRange(8, 48)
        self.font_spin.setValue(14)
        self.font_spin.valueChanged.connect(self.apply_font_size)
        top_row.addWidget(self.font_label)
        top_row.addWidget(self.font_spin)

        # 单双周切换按钮
        self.week_button = QPushButton('单周')
        self.week_button.clicked.connect(self.toggle_week_type)
        top_row.addWidget(self.week_button)
        
        # 保存设置
        self.save_button = QPushButton('保存设置')
        self.save_button.clicked.connect(self.save_settings_to_excel)
        top_row.addWidget(self.save_button)
        top_row.addStretch(1)
        layout.addLayout(top_row)

        # 顶部时间与控制行
        time_row = QHBoxLayout()
        self.date_label = QLabel()
        self.time_label = QLabel()
        time_row.addWidget(self.date_label)
        time_row.addWidget(self.time_label)

        # 自定义时间控制
        self.use_custom_time = QCheckBox('使用自定义时间')
        self.time_edit = QTimeEdit()
        self.time_edit.setDisplayFormat('HH:mm:ss')
        self.time_edit.setTime(datetime.now().time())
        time_row.addWidget(self.use_custom_time)
        time_row.addWidget(self.time_edit)

        time_row.addStretch(1)
        layout.addLayout(time_row)

        # 当前文件路径
        self.current_file = None
        
        # 单双周状态：0=单周，1=双周
        self.week_type = 0
        self.last_week_update = None

        # 表格
        self.table = QTableWidget()
        self.table.setColumnCount(2)
        self.table.setHorizontalHeaderLabels(['时间', '课程'])
        layout.addWidget(self.table)

        # 定时器
        self.timer = QTimer()
        self.timer.timeout.connect(self.update_time)
        self.timer.start(1000)

        self.schedule_data = None
        self.duty_data = None
        # 先读取设置（包括单双周状态），再加载课程表
        self.load_settings_from_excel()
        self.load_schedule()
        self.update_time()

        # 允许窗口与表格可伸缩
        self.table.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)

    def select_file(self):
        start_dir = self.get_app_dir()
        file_name, _ = QFileDialog.getOpenFileName(
            self,
            "选择课程表文件",
            start_dir,
            "Excel Files (*.xlsx *.xls);;All Files (*)"
        )
        if file_name:
            self.current_file = self.resolve_path(file_name)
            self.load_schedule()

    def load_schedule(self):
        try:
            if not self.current_file:
                default_file = os.path.join(self.get_app_dir(), 'schedule.xlsx')
                if os.path.exists(default_file):
                    self.current_file = default_file
                else:
                    QMessageBox.warning(self, "提示", "未找到默认课程表文件，请选择文件")
                    return

            df = pd.read_excel(self.current_file)
            weekday = datetime.now().weekday()

            # 根据单双周选择列范围：单周从A列(0)开始，双周从O列(14)开始
            base_col = 0 if self.week_type == 0 else 14  # 单周从0，双周从14（O列）
            
            time_col = base_col + weekday * 2
            course_col = time_col + 1
            
            # 检查列索引是否越界
            if course_col >= len(df.columns):
                # 如果双周时列数不够，回退到单周
                if self.week_type == 1:
                    QMessageBox.warning(self, "提示", "Excel文件中没有双周数据，已切换为单周")
                    self.week_type = 0
                    base_col = 0
                    time_col = base_col + weekday * 2
                    course_col = time_col + 1
                else:
                    QMessageBox.warning(self, "错误", f"Excel文件列数不足，需要至少{course_col + 1}列")
                    return

            raw_times = df.iloc[1:, time_col].tolist()
            raw_courses = df.iloc[1:, course_col].tolist()
            class_rows = []
            duty_rows = []
            mode = 'class'
            for t, c in zip(raw_times, raw_courses):
                t_str = '' if pd.isna(t) else str(t).strip()
                c_str = '' if pd.isna(c) else str(c).strip()
                token = t_str or c_str
                if token == 'END_C':
                    mode = 'duty'
                    continue
                if token == 'END':
                    break
                if mode == 'class':
                    class_rows.append((t_str, c_str))
                else:
                    duty_rows.append((t_str, c_str))

            times = [t for t, _ in class_rows]
            courses = [c for _, c in class_rows]

            self.table.setRowCount(len(times) + len(duty_rows))
            row_index = 0
            for t, c in zip(times, courses):
                self.table.setItem(row_index, 0, QTableWidgetItem(str(t)))
                self.table.setItem(row_index, 1, QTableWidgetItem(str(c)))
                row_index += 1
            for label, name in duty_rows:
                self.table.setItem(row_index, 0, QTableWidgetItem(str(label)))
                self.table.setItem(row_index, 1, QTableWidgetItem(str(name)))
                row_index += 1

            self.adjust_table_to_contents()
            self.apply_saved_column_widths()
            self.schedule_data = list(zip(times, courses))
            self.duty_data = duty_rows
        except Exception as e:
            print(f"加载课程表出错: {e}")

    def update_time(self):
        now_dt = datetime.now()
        if self.use_custom_time.isChecked():
            qt = self.time_edit.time()
            now_dt = datetime.combine(date.today(), time(qt.hour(), qt.minute(), qt.second()))

        weekdays = ['一', '二', '三', '四', '五', '六', '日']
        self.date_label.setText(f'星期{weekdays[now_dt.weekday()]}')
        suffix = ' (模拟)' if self.use_custom_time.isChecked() else ''
        self.time_label.setText(now_dt.strftime('%H:%M:%S') + suffix)

        if self.schedule_data:
            self.highlight_next_class(now_dt)
            self.highlight_duty_period(now_dt)

    def highlight_next_class(self, current_time):
        now_time = current_time.time()
        current_class_index = -1
        next_class_index = -1

        def parse_hhmm_to_time(hhmm: str):
            hhmm = hhmm.strip()
            try:
                parts = hhmm.split(':')
                if len(parts) != 2:
                    return None
                h = parts[0].zfill(2)
                m = parts[1].zfill(2)
                return datetime.strptime(f"{h}:{m}", '%H:%M').time()
            except Exception:
                return None

        for i, (time_str, _) in enumerate(self.schedule_data):
            if pd.isna(time_str):
                continue
            try:
                start_end = str(time_str).split('-')
                if len(start_end) != 2:
                    continue
                start_t = parse_hhmm_to_time(start_end[0])
                end_t = parse_hhmm_to_time(start_end[1])
                if not start_t or not end_t:
                    continue

                if start_t <= now_time < end_t:
                    current_class_index = i
                    break

                if next_class_index == -1 and start_t > now_time:
                    next_class_index = i
            except Exception:
                continue

        for i in range(self.table.rowCount()):
            for j in range(2):
                item = self.table.item(i, j)
                if item:
                    item.setBackground(QColor(255, 255, 255))

        target_index = current_class_index if current_class_index != -1 else next_class_index
        if target_index != -1:
            for j in range(2):
                item = self.table.item(target_index, j)
                if item:
                    item.setBackground(QColor(255, 255, 0))

    def highlight_duty_period(self, current_time):
        if not self.duty_data:
            return
        hour = current_time.hour
        highlight_prefix = None
        if 5 <= hour < 12:
            highlight_prefix = '上午'
        elif 12 <= hour < 22:
            highlight_prefix = '下午'
        if not highlight_prefix:
            return
        class_count = len(self.schedule_data) if self.schedule_data else 0
        for idx, (label, _) in enumerate(self.duty_data):
            if isinstance(label, str) and label.startswith(highlight_prefix):
                table_row = class_count + idx
                for j in range(2):
                    item = self.table.item(table_row, j)
                    if item:
                        item.setBackground(QColor(255, 255, 0))

    # ------------------ 设置读写 ------------------
    def load_settings_from_excel(self):
        if not self.current_file:
            return
        try:
            if not self.current_file.lower().endswith('.xlsx'):
                return
            wb = load_workbook(self.current_file)
            ws = wb.active
            settings = {}
            for row in range(30, 401):
                key = ws[f'A{row}'].value
                val = ws[f'B{row}'].value
                if key is None and val is None:
                    continue
                if isinstance(key, str):
                    settings[key.strip()] = val

            # 保存到实例，供列宽恢复使用
            self._loaded_settings = settings

            if 'font_size' in settings and settings['font_size']:
                try:
                    size = int(settings['font_size'])
                    if 8 <= size <= 48:
                        self.font_spin.setValue(size)
                        self.apply_font_size()
                except Exception:
                    pass

            if 'window_width' in settings and 'window_height' in settings:
                try:
                    w = int(settings['window_width'])
                    h = int(settings['window_height'])
                    if w > 200 and h > 200:
                        self.resize(w, h)
                except Exception:
                    pass
            
            # 读取单双周状态和最后更新时间
            if 'week_type' in settings and settings['week_type'] is not None:
                try:
                    self.week_type = int(settings['week_type']) % 2
                except Exception:
                    self.week_type = 0
            
            if 'last_week_update' in settings and settings['last_week_update']:
                try:
                    # 存储为字符串格式：YYYY-MM-DD HH:MM:SS
                    update_str = str(settings['last_week_update'])
                    self.last_week_update = datetime.strptime(update_str, '%Y-%m-%d %H:%M:%S')
                except Exception:
                    # 如果解析失败，使用当前时间
                    self.last_week_update = datetime.now()
            else:
                # 如果没有记录，初始化为当前时间
                self.last_week_update = datetime.now()
            
            # 根据时间差自动切换单双周
            self.calculate_auto_week_switch()
            self.update_week_button_text()
            
            # 如有 excel_path，支持相对路径（相对于应用目录）
            if 'excel_path' in settings and isinstance(settings['excel_path'], str):
                p = self.resolve_path(settings['excel_path'])
                if os.path.isfile(p) and p.lower().endswith(('.xlsx', '.xls')):
                    if self.current_file != p:
                        self.current_file = p
                        self.load_schedule()
        except Exception as e:
            print(f"读取设置失败: {e}")

    def save_settings_to_excel(self, show_message=True):
        if not self.current_file:
            if show_message:
                QMessageBox.information(self, '提示', '请先选择课程表文件再保存设置')
            return
        if not self.current_file.lower().endswith('.xlsx'):
            if show_message:
                QMessageBox.warning(self, '提示', '仅支持将设置写入 .xlsx 文件')
            return
        try:
            wb = load_workbook(self.current_file)
            ws = wb.active
            kv = {
                'font_size': self.font_spin.value(),
                'window_width': self.width(),
                'window_height': self.height(),
                'excel_path': self.current_file or '',
                'col_width_0': self.table.columnWidth(0) if self.table.columnCount() > 0 else 0,
                'col_width_1': self.table.columnWidth(1) if self.table.columnCount() > 1 else 0,
                'week_type': self.week_type,
                'last_week_update': self.last_week_update.strftime('%Y-%m-%d %H:%M:%S') if self.last_week_update else datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            }
            key_to_row = {}
            first_empty = None
            for row in range(30, 401):
                key_cell = ws[f'A{row}']
                val_cell = ws[f'B{row}']
                if key_cell.value is None and val_cell.value is None and first_empty is None:
                    first_empty = row
                if isinstance(key_cell.value, str):
                    key_to_row[key_cell.value.strip()] = row
            for key, val in kv.items():
                if key in key_to_row:
                    r = key_to_row[key]
                else:
                    r = first_empty if first_empty is not None else 30
                    while ws[f'A{r}'].value is not None or ws[f'B{r}'].value is not None:
                        r += 1
                ws[f'A{r}'] = key
                ws[f'B{r}'] = val
            wb.save(self.current_file)
            if show_message:
                QMessageBox.information(self, '提示', '设置已保存到 Excel')
        except Exception as e:
            QMessageBox.critical(self, '错误', f'保存设置失败: {e}')

    def apply_font_size(self):
        size = self.font_spin.value()
        font = self.table.font()
        font.setPointSize(size)
        self.table.setFont(font)
        f2 = self.date_label.font()
        f2.setPointSize(size)
        self.date_label.setFont(f2)
        f3 = self.time_label.font()
        f3.setPointSize(size)
        self.time_label.setFont(f3)
        self.adjust_table_to_contents()
        self.apply_saved_column_widths()

    def closeEvent(self, event):
        try:
            if self.current_file and self.current_file.lower().endswith('.xlsx'):
                self.save_settings_to_excel(show_message=False)
        finally:
            super().closeEvent(event)

    def apply_saved_column_widths(self):
        settings = getattr(self, '_loaded_settings', {}) or {}
        try:
            w0 = int(settings.get('col_width_0')) if settings.get('col_width_0') is not None else None
            w1 = int(settings.get('col_width_1')) if settings.get('col_width_1') is not None else None
            if w0 and self.table.columnCount() > 0:
                self.table.setColumnWidth(0, max(10, w0))
            if w1 and self.table.columnCount() > 1:
                self.table.setColumnWidth(1, max(10, w1))
        except Exception:
            pass

    def adjust_table_to_contents(self):
        self.table.resizeColumnsToContents()
        self.table.resizeRowsToContents()
        self.table.setSizeAdjustPolicy(QAbstractScrollArea.SizeAdjustPolicy.AdjustToContents)
        self.table.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.table.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.table.updateGeometry()
    
    def toggle_week_type(self):
        """手动切换单双周"""
        self.week_type = 1 - self.week_type  # 0变1，1变0
        self.last_week_update = datetime.now()
        self.update_week_button_text()
        self.load_schedule()
        # 自动保存
        if self.current_file and self.current_file.lower().endswith('.xlsx'):
            self.save_settings_to_excel(show_message=False)
    
    def update_week_button_text(self):
        """更新单双周按钮文本"""
        self.week_button.setText('单周' if self.week_type == 0 else '双周')
    
    def calculate_auto_week_switch(self):
        """根据时间差自动计算并切换单双周：计算两个周一之间的差值"""
        if self.last_week_update is None:
            return
        
        try:
            from datetime import timedelta
            now = datetime.now()
            
            # 计算当前时间所在周的周一
            now_monday = now.date() - timedelta(days=now.weekday())
            
            # 计算最后更新时间所在周的周一
            last_monday = self.last_week_update.date() - timedelta(days=self.last_week_update.weekday())
            
            # 计算两个周一之间的天数差
            delta = now_monday - last_monday
            days = delta.days
            
            # 每7天切换一次，计算切换次数
            weeks_passed = days // 7
            if weeks_passed > 0:
                # 切换相应次数（奇数次切换，偶数次不变）
                self.week_type = (self.week_type + weeks_passed) % 2
                self.last_week_update = now
                self.update_week_button_text()
        except Exception as e:
            print(f"自动切换单双周失败: {e}")

    # ------------------ 路径工具 ------------------
    def get_app_dir(self):
        try:
            # 冻结（打包）后：返回可执行文件所在目录；脚本运行：返回源文件目录
            if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
                return os.path.dirname(sys.executable)
            return os.path.dirname(os.path.abspath(__file__))
        except Exception:
            return os.getcwd()

    def resolve_path(self, path_str):
        try:
            if not path_str:
                return path_str
            if os.path.isabs(path_str):
                return os.path.normpath(path_str)
            return os.path.normpath(os.path.join(self.get_app_dir(), path_str))
        except Exception:
            return path_str


def main():
    app = QApplication(sys.argv)
    window = ScheduleWindowTest()
    window.show()
    sys.exit(app.exec())


if __name__ == '__main__':
    main()


